#!/usr/bin/env python3
"""Pin and summarize official TBEP/EPCHC long-term water-quality data.

This analysis uses the immutable Results_Updated.xls snapshot from tbep-tech/wq-static.
It follows the station-to-bay mapping and Middle Tampa Bay subsegment weights declared
in tbeptools. Monthly station means are formed first, then bay-segment means, then annual
means. This prevents stations/months with more samples from dominating annual values.

The output is a mechanism-screening covariate table. It does not establish causality.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from urllib.request import Request, urlopen

import numpy as np
import pandas as pd
import io
import openpyxl

WQ_COMMIT = "00aa86030f9245fe0318c186e7137798684dce74"
WQ_URL = (
    "https://raw.githubusercontent.com/tbep-tech/wq-static/"
    + WQ_COMMIT
    + "/data-raw/Results_Updated.xls"
)
WQ_SIZE = 61747627
WQ_BLOB = "b681a270e5cddb585fe5f7e03a853aa313bf0831"
SHEET = "RWMDataSpreadsheet"

STATION_SEGMENT = {}
for segment, stations in {
    "OTB": [36, 38, 40, 41, 46, 47, 50, 51, 60, 63, 64, 65, 66, 67, 68],
    "HB": [6, 7, 8, 44, 52, 55, 70, 71, 73, 80],
    "MTB": [9, 11, 81, 84, 13, 14, 32, 33, 16, 19, 28, 82],
    "LTB": [23, 24, 25, 90, 91, 92, 93, 95],
}.items():
    for station in stations:
        STATION_SEGMENT[station] = segment

MTB_GROUP = {}
for group, stations in {
    "MT1": [9, 11, 81, 84],
    "MT2": [13, 14, 32, 33],
    "MT3": [16, 19, 28, 82],
}.items():
    for station in stations:
        MTB_GROUP[station] = group

MTB_WEIGHT = {"MT1": 2108.7, "MT2": 1041.9, "MT3": 974.6}
MTB_WEIGHT_TOTAL = 4125.2

RAW_COLUMNS = {
    "station": "StationNumber",
    "sample_time": "SampleTime",
    "total_depth": "TotalDepth",
    "secchi": "SecchiDepth",
    "secchi_q": "Secchi_Q",
    "chlorophyll": "Chlorophyll_a",
    "total_nitrogen": "Total_Nitrogen",
    "salinity": "Sal-T",
    "temperature": "TempWater-T",
    "turbidity": "Turbidity",
}
METRICS = ["salinity", "temperature", "chlorophyll", "total_nitrogen", "secchi", "turbidity"]


def git_blob_sha1(data: bytes) -> str:
    return hashlib.sha1(f"blob {len(data)}\0".encode() + data).hexdigest()


def fetch() -> bytes:
    request = Request(WQ_URL, headers={"Accept-Encoding": "identity", "User-Agent": "Tampa-memory-study/1.0"})
    with urlopen(request, timeout=180) as response:
        data = response.read(WQ_SIZE + 1)
    if len(data) != WQ_SIZE:
        raise RuntimeError(f"water-quality source-size drift: {len(data)}")
    observed = git_blob_sha1(data)
    if observed != WQ_BLOB:
        raise RuntimeError(f"water-quality Git-blob drift: {observed}")
    return data


def as_float(value):
    if value in ("", None):
        return math.nan
    try:
        out = float(value)
    except (TypeError, ValueError):
        return math.nan
    return out if math.isfinite(out) else math.nan


def as_datetime(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        value = value.strip()
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%m/%d/%Y %H:%M", "%m/%d/%Y"):
            try:
                return datetime.strptime(value, fmt)
            except ValueError:
                pass
    return None


def load_selected(data: bytes) -> tuple[pd.DataFrame, dict]:
    # The upstream file is named .xls but is an OpenXML/XLSX ZIP container.
    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if SHEET not in book.sheetnames:
        raise RuntimeError(f"missing sheet: {SHEET}")
    sheet = book[SHEET]
    first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    header = [str(value).strip() if value is not None else "" for value in first]
    index = {name: i for i, name in enumerate(header)}
    missing = sorted(set(RAW_COLUMNS.values()) - set(index))
    if missing:
        raise RuntimeError(f"missing expected water-quality columns: {missing}")

    rows = []
    mapped_raw_rows = 0
    physical_rows = 1
    for values in sheet.iter_rows(min_row=2, values_only=True):
        physical_rows += 1
        station_value = as_float(values[index[RAW_COLUMNS["station"]]])
        if not math.isfinite(station_value):
            continue
        station = int(station_value)
        segment = STATION_SEGMENT.get(station)
        if segment is None:
            continue
        when = as_datetime(values[index[RAW_COLUMNS["sample_time"]]])
        if when is None or when.year < 1997 or when.year > 2025:
            continue

        mapped_raw_rows += 1
        total_depth = as_float(values[index[RAW_COLUMNS["total_depth"]]])
        secchi = as_float(values[index[RAW_COLUMNS["secchi"]]])
        secchi_q_raw = values[index[RAW_COLUMNS["secchi_q"]]]
        secchi_q = "" if secchi_q_raw is None else str(secchi_q_raw).strip()
        # tbeptools/read_formwq: VOB (">") and Secchi within 0.5 ft of bottom are unusable.
        if secchi_q == ">":
            secchi = math.nan
        if math.isfinite(total_depth) and math.isfinite(secchi):
            if (total_depth * 3.2809) - (secchi * 3.2809) < 0.5:
                secchi = math.nan

        rows.append(
            {
                "station": station,
                "segment": segment,
                "mtb_group": MTB_GROUP.get(station),
                "year": when.year,
                "month": when.month,
                "salinity": as_float(values[index[RAW_COLUMNS["salinity"]]]),
                "temperature": as_float(values[index[RAW_COLUMNS["temperature"]]]),
                "chlorophyll": as_float(values[index[RAW_COLUMNS["chlorophyll"]]]),
                "total_nitrogen": as_float(values[index[RAW_COLUMNS["total_nitrogen"]]]),
                "secchi": secchi,
                "turbidity": as_float(values[index[RAW_COLUMNS["turbidity"]]]),
            }
        )

    frame = pd.DataFrame(rows)
    audit = {
        "physical_rows_including_header": int(physical_rows),
        "sheet_columns": int(len(header)),
        "mapped_1997_2025_raw_rows": int(mapped_raw_rows),
        "selected_columns": RAW_COLUMNS,
        "container_format": "xlsx_openxml_despite_xls_filename",
    }
    book.close()
    return frame, audit


def aggregate_monthly(frame: pd.DataFrame) -> pd.DataFrame:
    # One value per station-month first.
    station_month = (
        frame.groupby(["station", "segment", "mtb_group", "year", "month"], dropna=False)[METRICS]
        .mean()
        .reset_index()
    )

    rows = []
    for (segment, year, month), group in station_month.groupby(["segment", "year", "month"]):
        row = {"segment": segment, "year": int(year), "month": int(month)}
        for metric in METRICS:
            if segment == "MTB":
                subgroup_values = {}
                for subgroup, sub in group.groupby("mtb_group"):
                    values = sub[metric].dropna()
                    if len(values):
                        subgroup_values[subgroup] = float(values.mean())
                present_weight = sum(MTB_WEIGHT[g] for g in subgroup_values)
                row[metric] = (
                    sum(subgroup_values[g] * MTB_WEIGHT[g] for g in subgroup_values) / present_weight
                    if present_weight
                    else math.nan
                )
                row[f"{metric}_mtb_weight_fraction"] = present_weight / MTB_WEIGHT_TOTAL
            else:
                values = group[metric].dropna()
                row[metric] = float(values.mean()) if len(values) else math.nan
                row[f"{metric}_mtb_weight_fraction"] = math.nan
        row["station_count"] = int(group.station.nunique())
        rows.append(row)
    return pd.DataFrame(rows).sort_values(["segment", "year", "month"]).reset_index(drop=True)


def aggregate_annual(monthly: pd.DataFrame) -> pd.DataFrame:
    records = []
    for (segment, year), group in monthly.groupby(["segment", "year"]):
        row = {"segment": segment, "year": int(year), "months_present": int(group.month.nunique())}
        for metric in METRICS:
            values = group[metric].dropna()
            row[metric] = float(values.mean()) if len(values) else math.nan
            row[f"{metric}_months"] = int(len(values))
        records.append(row)
    return pd.DataFrame(records).sort_values(["segment", "year"]).reset_index(drop=True)


def pre2016_anomaly_table(annual: pd.DataFrame) -> pd.DataFrame:
    records = []
    for segment in ["OTB", "MTB"]:
        segment_frame = annual[annual.segment == segment]
        for metric in METRICS:
            baseline = segment_frame[(segment_frame.year >= 1997) & (segment_frame.year <= 2015)][metric].dropna()
            current = segment_frame.loc[segment_frame.year == 2016, metric]
            if len(baseline) < 8 or len(current) != 1 or not math.isfinite(float(current.iloc[0])):
                continue
            x = float(current.iloc[0])
            median = float(baseline.median())
            mad = float(np.median(np.abs(baseline.to_numpy() - median)))
            robust_z = (x - median) / (1.4826 * mad) if mad > 0 else math.nan
            percentile = float((baseline < x).mean())
            prev = segment_frame.loc[segment_frame.year == 2015, metric]
            nxt = segment_frame.loc[segment_frame.year == 2017, metric]
            records.append(
                {
                    "segment": segment,
                    "metric": metric,
                    "value_2016": x,
                    "pre2016_median": median,
                    "pre2016_mad": mad,
                    "robust_z_vs_1997_2015": robust_z,
                    "pre2016_percentile": percentile,
                    "value_2015": float(prev.iloc[0]) if len(prev) else math.nan,
                    "value_2017": float(nxt.iloc[0]) if len(nxt) else math.nan,
                    "delta_2016_minus_2015": x - float(prev.iloc[0]) if len(prev) else math.nan,
                    "delta_2017_minus_2016": float(nxt.iloc[0]) - x if len(nxt) else math.nan,
                }
            )
    return pd.DataFrame(records)


def main(outdir: Path) -> None:
    outdir.mkdir(parents=True, exist_ok=True)
    data = fetch()
    frame, audit = load_selected(data)
    monthly = aggregate_monthly(frame)
    annual = aggregate_annual(monthly)
    anomaly = pre2016_anomaly_table(annual)

    summary = {
        "source": {
            "repository": "tbep-tech/wq-static",
            "commit": WQ_COMMIT,
            "file": "data-raw/Results_Updated.xls",
            "size_bytes": WQ_SIZE,
            "git_blob_sha1": WQ_BLOB,
            "verified": True,
        },
        "parser_audit": audit,
        "years": [int(annual.year.min()), int(annual.year.max())],
        "segments": sorted(annual.segment.unique().tolist()),
        "annual_rows": int(len(annual)),
        "anomaly_2016": anomaly.replace({np.nan: None}).to_dict("records"),
        "claim_boundary": (
            "Response-independent environmental screening. Annual segment anomalies are "
            "descriptive and do not identify the cause of the Thalassia turnover pulse."
        ),
    }

    frame.to_csv(outdir / "water_quality_selected_raw.csv.gz", index=False, compression="gzip")
    monthly.to_csv(outdir / "water_quality_monthly.csv", index=False)
    annual.to_csv(outdir / "water_quality_annual.csv", index=False)
    anomaly.to_csv(outdir / "water_quality_2016_anomaly.csv", index=False)
    (outdir / "water_quality_summary.json").write_text(json.dumps(summary, indent=2, sort_keys=True))
    print(json.dumps(summary, indent=2, sort_keys=True))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, default=Path("results/generated"))
    args = parser.parse_args()
    main(args.out)
