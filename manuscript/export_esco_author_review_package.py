#!/usr/bin/env python3
"""Build an auditable, clearly non-final ESCO submission package from canonical Markdown.

This export DOES NOT update scientific results, generate author approvals,
assert new permits, choose a new forecast, or infer seagrass mortality.
Inputs must be exactly the verified source manuscript and its generated
supplement. External figure/table artifacts are required and pinned to the
scientific PR head in the companion workflow.

Requires pypandoc-binary, python-docx and libreoffice. PyMuPDF is used for
content and page-geometry QA. Executable both on GitHub Actions and locally.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
import zipfile

import fitz
import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MAIN = ROOT / "manuscript/TAMPA_ECOLOGY_ESCO_SUBMISSION_V1.md"
SUPP = ROOT / "manuscript/TAMPA_ECOLOGY_SUPPLEMENT_ASSEMBLED_V1.md"
BACK = ROOT / "manuscript/ESTUARIES_COASTS_BACK_MATTER_TEMPLATE_V1.md"
LETTER = ROOT / "manuscript/ESTUARIES_COASTS_COVER_LETTER_TEMPLATE_V1.md"
REPORTING = ROOT / "results/warning_reporting_metrics_v1.json"
MARKS = ROOT / "results/early_warning_fixed_mark_overlap_verified_v1.json"
TITLE = "Persistent occurrence conceals spatially heterogeneous degradation in a seagrass foundation species"
REQUIRED_MAIN_PHRASES = [
    "1,497", "1,480", "71", "2016", "Thalassia",
    "post-hoc", "not interpreted as demographic extinction",
    "AUC", "Braun", "Fig."
]
REQUIRED_SUPP_PHRASES = [
    "Supplementary Methods S1", "Supplementary Methods S14",
    "Fig. S1", "Fig. S7", "Table S1", "Table S12",
    "not demographic extinction"
]

def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for part in iter(lambda: f.read(65536), b""):
            h.update(part)
    return h.hexdigest()

def source_check():
    original = MAIN.read_text(encoding="utf-8")
    supplement = SUPP.read_text(encoding="utf-8")
    back = BACK.read_text(encoding="utf-8")
    letter = LETTER.read_text(encoding="utf-8")
    assert original.startswith("# " + TITLE), "Wrong/obsolete ESCO manuscript title"
    assert all(t in original for t in REQUIRED_MAIN_PHRASES), "Source manuscript missing previously verified claim"
    assert all(t in supplement for t in REQUIRED_SUPP_PHRASES), "Source supplement missing canonical captions or methods"
    assert "Statements and Declarations" in back, "Missing author-supplied back matter"
    assert "Author attestations" in letter, "Missing actual author-approval gate"
    assert len(original)>30000 and len(supplement)>20000
    # Only author-specific templated metadata is inserted; no hypothesis,
    # sample, reference or ecological conclusion is altered.
    return original,supplement,back,letter

def page_field(para):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    para._p.append(fld)

def set_continuous_line_numbers(sec):
    pr=sec._sectPr
    old=pr.findall(qn("w:lnNumType"))
    for node in old:pr.remove(node)
    ln=OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:start"), "1")
    ln.set(qn("w:distance"), "360")
    ln.set(qn("w:restart"), "continuous")
    pr.append(ln)

def style_doc(doc: Document, running_title: str):
    styles=doc.styles
    for name in ("Normal","Body Text","Heading 1","Heading 2","Heading 3","Title"):
        if name not in styles:continue
        s=styles[name]
        s.font.name="Times New Roman"
        s.font.color.rgb=RGBColor(29,43,54)
        if name in ("Normal","Body Text"):
            s.font.size=Pt(10)
            s.paragraph_format.space_after=Pt(5)
            s.paragraph_format.line_spacing=1.42
        elif name.startswith("Heading"):
            s.font.bold=True
            s.font.size=Pt({"Heading 1":12,"Heading 2":11,"Heading 3":10}[name])
            s.paragraph_format.space_before=Pt(11)
            s.paragraph_format.space_after=Pt(5)
            s.paragraph_format.keep_with_next=True
        else:
            s.font.size=Pt(15)
            s.font.bold=True
    for section in doc.sections:
        section.top_margin=Inches(.83)
        section.bottom_margin=Inches(.79)
        section.left_margin=Inches(.86)
        section.right_margin=Inches(.84)
        section.header_distance=Inches(.38)
        section.footer_distance=Inches(.42)
        set_continuous_line_numbers(section)
        header=section.header.paragraphs[0]
        header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        header.text="AUTHOR REVIEW  |  NOT SUBMITTED"
        for run in header.runs:
            run.font.name="Arial"
            run.font.size=Pt(7)
            run.font.color.rgb=RGBColor(112,123,136)
        footer=section.footer.paragraphs[0]
        footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run("Page ").font.size=Pt(8)
        page_field(footer)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:run.font.size=Pt(8)
    doc.core_properties.author="Author team (not supplied)"
    doc.core_properties.subject="Non-final source-backed ESCO author-review export"
    doc.core_properties.keywords="Tampa Bay; seagrass; source-derived manuscript; not submitted"

def convert_markdown(md: str, filename: Path, resource_dirs: list[Path], title: str):
    tmp=filename.with_suffix(".source.md")
    tmp.parent.mkdir(parents=True,exist_ok=True)
    tmp.write_text(md,encoding="utf-8")
    args=["--standalone","--reference-doc="+str(filename.parent/"reference_style.docx"),
          "--resource-path="+":".join(str(p) for p in resource_dirs)]
    pypandoc.convert_file(str(tmp),"docx",format="markdown+tex_math_dollars",
                          outputfile=str(filename),extra_args=args)
    doc=Document(str(filename))
    style_doc(doc,title)
    doc.save(str(filename))
    return filename

def clean_author_placeholders(block: str)->str:
    # Preserve missing metadata visibly instead of feeding angle-bracket
    # placeholders to Pandoc's HTML parser, which can swallow them.
    return re.sub(r"<([^<>\n]{2,160})>",lambda m: "**[AUTHOR INPUT REQUIRED: "+m.group(1)+"]**",block)

def main_markup(source: str, back: str)->str:
    first=source.split("\n",1)[0]
    assert first.startswith("# ")
    body=source[len(first):].lstrip()
    titlepage=(
      first+"\n\n"
      "**Non-final Original Article author-review proof — DO NOT SUBMIT AS IS**\n\n"
      "**Authors:** [AUTHOR INPUT REQUIRED: complete verified author list]\n\n"
      "**Affiliations:** [AUTHOR INPUT REQUIRED: complete institution details]\n\n"
      "**Corresponding author:** [AUTHOR INPUT REQUIRED: name and active email]\n\n"
      "**Submission date:** [AUTHOR INPUT REQUIRED: date of actual submission]\n\n"
      "---\n\n"
    )
    start=back.index("## Acknowledgements")
    declarations=clean_author_placeholders(back[start:].strip())
    return titlepage+body+"\n\n# Author and journal declarations — not yet approved\n\n"+declarations+"\n"

def supp_markup(source: str,figdir: Path)->str:
    md=(
      "# Supplementary Information — AUTHOR REVIEW PROOF\n\n"
      "**Title:** "+TITLE+"\n\n"
      "**Journal:** Estuaries and Coasts (proposed original article)\n\n"
      "**Authors/affiliations/correspondence:** [AUTHOR INPUT REQUIRED]\n\n"
      "**Status:** Not submitted. Figures S1–S7 are appended from source-head CI artwork; "
      "tables S1–S12 are supplied as separate machine-readable CSV Online Resource candidates.\n\n"
      "---\n\n"+source+"\n\n"
      "# Supplementary Figure Plates S1–S7\n\n"
      "These source-generated figures are reproduced for reviewer readability; "
      "their full captions and inferential boundaries appear above. "
      "Confirm the final figure labels and visual scale before submission.\n\n"
    )
    for n in range(1,8):
        p=figdir/f"FigS{n}_QA.png"
        assert p.is_file() and p.stat().st_size>5000, f"Missing source S{n} figure {p}"
        md+=f"## Figure S{n} artwork\n\n![Source-rendered Supplementary Figure S{n}](FigS{n}_QA.png){{ width=6.1in }}\n\n"
    return md

def supplementary_tables(out_dir: Path, original_csv: Path):
    t=out_dir/"OnlineResource_Tables"
    t.mkdir(parents=True,exist_ok=True)
    files=sorted(original_csv.glob("Table_S*.csv"))
    assert len(files)==9 and all(f"Table_S{i}_" in " ".join(p.name for p in files) for i in range(1,10)), \
        "Source table artifact lacks one of S1-S9"
    for p in files:
        with p.open(encoding="utf-8-sig",newline="") as h:
            assert len(list(csv.DictReader(h)))>0, "Empty table "+p.name
        shutil.copy2(p,t/p.name)
    metric=json.loads(REPORTING.read_text(encoding="utf-8"))
    marks=json.loads(MARKS.read_text(encoding="utf-8"))
    assert metric["schema"]=="tampa.warning_reporting_metrics_v1"
    assert marks["status"]=="COMPLETED_OBSERVATION_PROCESS_AUDIT"

    def write(name,cols,rows):
        dst=t/name
        with dst.open("w",newline="",encoding="utf-8") as h:
            writer=csv.DictWriter(h,fieldnames=cols)
            writer.writeheader()
            for row in rows:writer.writerow(row)
        return dst

    # S10: full source-ledger arms; no new model.
    labels={
        "space":"Space/time/effort",
        "space_quant":"Space/time/effort + quantity",
        "node":"L2-regularized node ID + space/time/effort",
        "node_quant":"Same regularized reference + quantity"}
    cols=["model","descriptor","pooled_auc","pooled_log_loss","macro_annual_log_loss","brier"]
    rows=[]
    for key,m in metric["metrics"].items():
        rows.append(dict(model=key,descriptor=labels[key],pooled_auc=m["auc"],
                         pooled_log_loss=m["log_loss"],
                         macro_annual_log_loss=m["macro_log_loss"],brier=m["brier"]))
    write("Table_S10_node_identity_predictive_comparisons.csv",cols,rows)

    cols=["partition","events","no_positive_mark_revisited","some_but_not_all_positive_marks_revisited",
          "all_positive_marks_revisited","at_least_one_positive_mark_revisited"]
    rows=[]
    for key,v in marks["loss_partitions"].items():
        rows.append(dict(partition=key,events=v["events"],
                         no_positive_mark_revisited=v["no_positive_mark_revisited"],
                         some_but_not_all_positive_marks_revisited=v["some_but_not_all_revisited"],
                         all_positive_marks_revisited=v["all_positive_marks_revisited"],
                         at_least_one_positive_mark_revisited=v["at_least_one_revisited"]))
    write("Table_S11_recorded_loss_mark_revisits.csv",cols,rows)

    cols=["annual_monitoring_capacity_fraction","sum_node_year_alerts",
          "site_reference_loss_events_captured","quantitative_loss_events_captured",
          "site_reference_revisited_loss_events_captured",
          "quantitative_revisited_loss_events_captured","source_positive_scored_losses",
          "n_changed_watchlist_members"]
    rows=[]
    for v in marks["capacity_decomposition"]:
        rows.append(dict(
           annual_monitoring_capacity_fraction=v["nominal_capacity"],
           sum_node_year_alerts=v["alerts_per_arm"],
           site_reference_loss_events_captured=v["captured_all_node"],
           quantitative_loss_events_captured=v["captured_all_node_quant"],
           site_reference_revisited_loss_events_captured=v["captured_revisited_node"],
           quantitative_revisited_loss_events_captured=v["captured_revisited_node_quant"],
           source_positive_scored_losses=marks["registry"]["scored_losses"],
           n_changed_watchlist_members=v["selection_changes"]))
    write("Table_S12_capacity_and_revisit_capture.csv",cols,rows)
    assert len(list(t.glob("Table_S*.csv")))==12
    return t

def convert_pdf_with_lo(doc: Path, destination: Path):
    if not shutil.which("libreoffice"):
        raise RuntimeError("LibreOffice required for PDF/proof rendering on runner")
    cwd=destination.parent
    tmp=cwd/"pdf_render"
    tmp.mkdir(parents=True,exist_ok=True)
    cp=subprocess.run(["libreoffice","-env:UserInstallation=file:///tmp/lo_esco_author_review",
                       "--headless","--convert-to","pdf:writer_pdf_Export",
                       "--outdir",str(tmp),str(doc)],
                      stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True,timeout=200)
    made=tmp/(doc.stem+".pdf")
    if cp.returncode or not made.is_file():
        raise RuntimeError(f"LibreOffice PDF failed: {cp.returncode}; {cp.stdout[:300]} {cp.stderr[:400]}")
    shutil.copy2(made,destination)
    return destination

def pdf_review(pdf: Path, needed: list[str], min_pages: int, min_embedded_images: int=0):
    d=fitz.open(str(pdf))
    assert d.page_count>=min_pages,(pdf.name,d.page_count)
    text="\n".join(p.get_text(sort=True) for p in d)
    for needle in needed:
        assert needle.lower() in text.lower(), f"{pdf.name}: missing expected {needle!r}"
    assert "AUTHOR REVIEW" in text or "Author Review" in text or "author-review" in text
    image_count=0
    bbox_failures=[]
    for i,p in enumerate(d):
        image_count+=len(p.get_images(full=True))
        for b in p.get_text("blocks"):
            if not b or len(b)<4 or not isinstance(b[4],str):continue
            x0,y0,x1,y1=b[:4]
            if x0 < -2 or y0 < -2 or x1 > p.rect.width+2 or y1 > p.rect.height+2:
                bbox_failures.append([i+1,round(x0,1),round(y0,1),round(x1,1),round(y1,1)])
    assert not bbox_failures[:6],(pdf.name,"clipped",bbox_failures[:6])
    assert image_count>=min_embedded_images,(pdf.name,"figures",image_count)
    return dict(file=pdf.name,pages=d.page_count,
                image_references=image_count,bbox_clipping_found=False,
                source_keywords_checked=len(needed))

def docx_review(doc: Path, need: list[str]):
    d=Document(str(doc))
    text="\n".join(p.text for p in d.paragraphs)
    assert len(text)>1000,doc.name
    for x in need:
        assert x.lower() in text.lower(),(doc.name,x)
    styles=[sec._sectPr.find(qn("w:lnNumType")) is not None for sec in d.sections]
    assert all(styles),"Continuous line-number setting missing from Word section"
    return {"file":doc.name,"paragraphs":len(d.paragraphs),
            "continuous_line_number_sections":sum(styles),"checks":len(need)}

def snapshot_zip(artifacts: Path, archive: Path):
    with zipfile.ZipFile(archive,"w",compression=zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for file in sorted(artifacts.rglob("*")):
            if file.is_file() and file!=archive:
                z.write(file,file.relative_to(artifacts))
    return archive

def main():
    p=argparse.ArgumentParser()
    p.add_argument("--out",type=Path,default=ROOT/"results/generated_esco_author_review")
    p.add_argument("--main-figures",type=Path,default=ROOT/"results/generated_esco_figures")
    p.add_argument("--supp-figures",type=Path,default=ROOT/"results/generated_supplement_figures_rendered")
    p.add_argument("--supp-tables",type=Path,default=ROOT/"results/generated_supplement_tables")
    a=p.parse_args()
    out=a.out.resolve()
    out.mkdir(parents=True,exist_ok=True)
    main_src,supp_src,back,letter=source_check()
    assets=out/"source_figures"
    assets.mkdir(exist_ok=True)
    for i in range(1,5):
        for ext in ("eps","tif"):
            f=a.main_figures/f"Fig{i}.{ext}"
            assert f.is_file() and f.stat().st_size>3000, f"Missing verified main figure: {f}"
            shutil.copy2(f,assets/f.name)
    for i in range(1,8):
        for ext in ("eps","tif"):
            f=a.supp_figures/f"FigS{i}.{ext}"
            assert f.is_file() and f.stat().st_size>3000, f"Missing verified SI figure: {f}"
            shutil.copy2(f,assets/f.name)

    template=out/"reference_style.docx"
    style=Document()
    style_doc(style,TITLE)
    style.save(str(template))
    tables=supplementary_tables(out,a.supp_tables)

    main_doc=out/"Tampa_Main_AuthorReview.docx"
    si_doc=out/"Tampa_Supplement_AuthorReview.docx"
    letter_doc=out/"Tampa_CoverLetter_TEMPLATE.docx"
    convert_markdown(main_markup(main_src,back),main_doc,[ROOT/"manuscript"],TITLE)
    convert_markdown(supp_markup(supp_src,a.supp_figures),si_doc,[a.supp_figures,ROOT/"manuscript"],"Supplementary Information")
    convert_markdown(
      "**AUTHOR APPROVALS AND CORRESPONDING AUTHOR DETAILS ARE NOT CONFIRMED. THIS IS A TEMPLATE.**\n\n"+letter,
      letter_doc,[ROOT/"manuscript"],"Cover Letter - Author Input Required"
    )

    main_pdf=convert_pdf_with_lo(main_doc,out/"Tampa_Main_Proof_NotSubmitted.pdf")
    supplement_pdf=convert_pdf_with_lo(si_doc,out/"Tampa_Supplement_AuthorReview.pdf")
    letter_pdf=convert_pdf_with_lo(letter_doc,out/"Tampa_CoverLetter_Proof_NotSubmitted.pdf")

    validation={
      "status":"AUTHOR_REVIEW_ARTIFACTS_ONLY_NOT_SUBMITTED",
      "article_title":TITLE,
      "target_journal":"Estuaries and Coasts",
      "source_inputs":{
        str(x.relative_to(ROOT)):sha256(x) for x in
        (MAIN,SUPP,BACK,LETTER,REPORTING,MARKS)},
      "provenance":{
        "main_figure_generator":"figures/render_esco_figures.py",
        "supplement_figure_artifact_run":39880446814,
        "supplement_table_artifact_run":39880446784,
        "source_pr":"https://github.com/zuizui0223/Tampa/pull/155",
        "not_frozen_submission_git_release":True},
      "docx_checks":[
          docx_review(main_doc,[TITLE,"1,497","1,480","2016","Figure captions","References"]),
          docx_review(si_doc,["Supplementary Information","Supplementary Methods S14","Table S12"]),
          docx_review(letter_doc,["Estuaries and Coasts","Author attestations"])],
      "render_checks":[
          pdf_review(main_pdf,["Persistent occurrence","Abstract","Methods","Results","References"],10),
          pdf_review(supplement_pdf,["Supplementary Information","Supplementary Methods S14","Table S12","Figure S7 artwork"],10,7),
          pdf_review(letter_pdf,["Estuaries and Coasts","corresponding author"],1)],
      "exported_tables_count":len(list(tables.glob("Table_S*.csv"))),
      "author_approval_blockers":[
        "final author list, affiliations, corresponding email and signed contributions",
        "source-data permissions and journal-required declarations",
        "corresponding author's cover-letter attestations; coauthor approvals",
        "visual proof of all Word/PDF pages and final figure legends",
        "journal-specific online resource filenames, manuscript formatting and actual submission",
        "freeze release commit/tag and source code archive DOI if created"
      ]
    }
    important=[
        main_doc,si_doc,supplement_pdf,letter_doc,main_pdf,letter_pdf,
        *sorted(assets.iterdir()),*sorted(tables.iterdir())
    ]
    validation["files"]={str(f.relative_to(out)):{
        "bytes":f.stat().st_size,"sha256":sha256(f)} for f in important}
    report=out/"SOURCE_AND_LAYOUT_AUDIT.json"
    report.write_text(json.dumps(validation,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    (out/"AUTHOR_ACTIONS_BEFORE_SUBMISSION.md").write_text(
       "# Tampa seagrass Original Article — author-review export\n\n"
       "**This package was built from source-verified repository inputs and has NOT been submitted.**\n\n"
       "The ESCO main DOCX contains author placeholders and source-derived declarations; the supplementary PDF contains the S1–S14 methods text and source figure plates S1–S7. "
       "Source-authenticated machine-readable Tables S1–S12 and TIFF/EPS images are provided alongside. "
       "Do not treat existence of an export as contributor approval or final journal permission.\n\n"
       "## Human checks still necessary\n\n"+
       "\n".join(f"- [ ] {x}" for x in validation["author_approval_blockers"])+
       "\n\nInspect the full Word, PDF and all image files before upload. Confirm the live journal submission requirements. "
       "Do not automatically attest that the manuscript is original and not under consideration elsewhere.\n",
       encoding="utf-8")
    # Do not ship intermediate reference_docx, source Markdown or PDF-render folder.
    template.unlink(missing_ok=True)
    for f in (out/"Tampa_Main_AuthorReview.source.md",
              out/"Tampa_Supplement_AuthorReview.source.md",
              out/"Tampa_CoverLetter_TEMPLATE.source.md"):
        f.unlink(missing_ok=True)
    shutil.rmtree(out/"pdf_render",ignore_errors=True)
    archive=out.parent/"Tampa_ESCO_AUTHOR_REVIEW_PACKAGE.zip"
    snapshot_zip(out,archive)
    print(json.dumps({"status":validation["status"],
        "pdf_checks":validation["render_checks"],
        "docx_checks":validation["docx_checks"],
        "source_csv_table_count":validation["exported_tables_count"],
        "file_count":len(validation["files"]),
        "archive":str(archive),"archive_bytes":archive.stat().st_size,
        "archive_sha256":sha256(archive)},indent=2))

if __name__=="__main__":
    main()
