#!/usr/bin/env python3
"""Response-blind metadata qualification for Florida Keys seagrass v1."""
from __future__ import annotations
import hashlib, io, json, re
from pathlib import Path
from urllib.request import Request, urlopen
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
C = json.loads((ROOT / "metadata_contract.json").read_text())
OUT = ROOT / "metadata_result.json"

def canon(x):
    return hashlib.sha256(json.dumps(x, sort_keys=True, separators=(",",":"), ensure_ascii=True).encode()).hexdigest()

def fetch(url, ua):
    req = Request(url, headers={"Accept-Encoding":"identity","User-Agent":ua})
    with urlopen(req, timeout=60) as r:
        if r.status != 200:
            raise RuntimeError(f"http_{r.status}:{url}")
        return r.read()

def norm(x):
    return re.sub(r"\s+"," ",str(x or "").strip())

def main():
    out = {
        "schema":"tampa.florida_keys_external_early_warning_v1.metadata_result",
        "status":None,
        "biological_data_downloads":0,
        "biological_response_rows_opened":0,
        "biological_response_values_opened":False
    }
    try:
        page = fetch(C["source"]["program_page"], "Tampa-FK-Seagrass-Metadata/1.0").decode("utf-8", errors="strict")
        required_page = [
            "Florida Keys National Marine Sanctuary Seagrass Monitoring Project",
            "1995",
            "Fixed transect",
            "Braun-Blanquet",
            "SAV - 296.zip"
        ]
        missing_page = [x for x in required_page if x not in page]
        if missing_page:
            raise RuntimeError("program_page_drift:"+"|".join(missing_page))

        xlsx = fetch(C["source"]["public_metadata_workbook"], "Tampa-FK-Seagrass-Metadata/1.0")
        wb = load_workbook(io.BytesIO(xlsx), read_only=True, data_only=True)
        terms = [
            "percent cover","braun","species","taxon","station","site","transect",
            "sample date","sampling date","year","absence","absent","zero","0",
            "missing","not observed","submerged aquatic vegetation"
        ]
        hits = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                vals = [norm(c.value) for c in row]
                joined = " | ".join(v for v in vals if v)
                low = joined.lower()
                if joined and any(t in low for t in terms):
                    hits.append({"sheet":ws.title,"row":row[0].row,"text":joined[:1200]})
                    if len(hits) >= 250:
                        break
            if len(hits) >= 250:
                break

        safe = {
            "program_page_sha256": hashlib.sha256(page.encode("utf-8")).hexdigest(),
            "metadata_xlsx_sha256": hashlib.sha256(xlsx).hexdigest(),
            "metadata_xlsx_size": len(xlsx),
            "sheet_names": wb.sheetnames,
            "metadata_hits": hits
        }
        out["safe_metadata"] = safe
        out["safe_metadata_fingerprint"] = canon(safe)
        out["status"] = "metadata_inventory_complete"
        out["next_gate"] = "inspect metadata-only hits and either freeze exact semantics or stop pre-response"
    except Exception as exc:
        out["status"] = "terminal_pre_response_metadata_stop"
        out["reason"] = str(exc)
        out["next_gate"] = "none"
    out["fingerprint"] = canon(out)
    OUT.write_text(json.dumps(out, indent=2, sort_keys=True) + "\n")
    print(json.dumps(out, indent=2, sort_keys=True))

if __name__ == "__main__":
    main()
