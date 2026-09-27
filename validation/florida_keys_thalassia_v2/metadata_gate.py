#!/usr/bin/env python3
"""Metadata-only SEACAR schema gate for Florida Keys Thalassia v2."""
from __future__ import annotations
import hashlib, io, json, re
from pathlib import Path
from urllib.request import Request, urlopen
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parent
C=json.loads((ROOT/"metadata_contract.json").read_text())
OUT=ROOT/"metadata_result.json"

def canon(x):
    return hashlib.sha256(json.dumps(x,sort_keys=True,separators=(",",":"),ensure_ascii=True).encode()).hexdigest()

def norm(x):
    return re.sub(r"\s+"," ",str(x or "").strip())

def main():
    out={
        "schema":"tampa.florida_keys_external_early_warning_v2.metadata_result",
        "status":None,
        "biological_data_downloads":0,
        "biological_response_rows_opened":0,
        "biological_response_values_opened":False
    }
    try:
        url=C["metadata_source"]["url"]
        req=Request(url,headers={"Accept-Encoding":"identity","User-Agent":"Tampa-FK-Thalassia-v2-Metadata/1.0"})
        with urlopen(req,timeout=60) as resp:
            if resp.status!=200:
                raise RuntimeError(f"metadata_http_{resp.status}")
            raw=resp.read(20_000_001)
        if len(raw)>20_000_000:
            raise RuntimeError("metadata_workbook_too_large")
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        keywords=[
            "program","programid","program id","station","site","location","sample",
            "date","year","species","scientific","taxon","percent cover","cover",
            "submerged aquatic vegetation","sav","absence","absent","zero","missing","null"
        ]
        hits=[]
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                vals=[norm(c.value) for c in row]
                joined=" | ".join(v for v in vals if v)
                low=joined.lower()
                if joined and any(k in low for k in keywords):
                    hits.append({"sheet":ws.title,"row":row[0].row,"text":joined[:1600]})
                    if len(hits)>=500:
                        break
            if len(hits)>=500:
                break
        safe={
            "workbook_sha256":hashlib.sha256(raw).hexdigest(),
            "workbook_size":len(raw),
            "sheet_names":wb.sheetnames,
            "metadata_hits":hits
        }
        out["safe_metadata"]=safe
        out["safe_metadata_fingerprint"]=canon(safe)
        out["status"]="metadata_inventory_complete"
        out["next_gate"]="interpret only metadata hits against frozen required semantics; no biological data access"
    except Exception as exc:
        out["status"]="terminal_pre_response_metadata_stop"
        out["reason"]=str(exc)
        out["next_gate"]="none"
    out["fingerprint"]=canon(out)
    OUT.write_text(json.dumps(out,indent=2,sort_keys=True)+"\n")
    print(json.dumps(out,indent=2,sort_keys=True))

if __name__=="__main__":
    main()
